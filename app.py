import streamlit as st
import pandas as pd
import numpy as np
import docx
from docxtpl import DocxTemplate
from docx.oxml import OxmlElement
import io
import os
import base64
import re
from datetime import datetime
import gspread
from google.oauth2.service_account import Credentials



# --- KONFIGURASI TAMPILAN WEB ---

st.set_page_config(page_title="SIBRIO - Generator BRS", page_icon="logobps.png", layout="wide")



# --- INJEKSI CUSTOM CSS ---

st.markdown("""

<style>

@import url('https://fonts.googleapis.com/css2?family=Poppins:wght@400;600;800;900&display=swap');



/* Mengatur font global dan background hitam */

.stApp {

    background-color: #0E1117;

    font-family: 'Poppins', sans-serif;

    color: white;

}



/* Menyembunyikan padding atas bawaan Streamlit */

.block-container {

    padding-top: 2rem !important;

}



/* Container untuk Logo & Judul Tengah */

.header-container {

    display: flex;

    justify-content: center;

    align-items: center;

    gap: 20px;

    margin-top: 10px;

}



/* Main Title SIBRIO */

.main-title {

    color: #FF5722;

    font-size: 65px;

    font-weight: 900;

    letter-spacing: 3px;

    margin: 0;

    line-height: 1.1;

}

.main-subtitle {

    text-align: center;

    color: #E0E0E0;

    font-size: 16px;

    font-weight: 400;

    margin-bottom: 40px;

    line-height: 1.4;

    margin-top: 10px;

}



/* Label Custom Rata Tengah */

.custom-label {

    text-align: center;

    font-weight: 600;

    font-size: 14px;

    color: #FAFAFA;

    margin-bottom: 5px;

}



/* Styling Tombol Generate */

div.stButton > button:first-child {

    background-color: #FF5722 !important;

    color: white !important;

    border-radius: 8px !important;

    border: none !important;

    height: 50px !important;

    font-size: 16px !important;

    font-weight: 600 !important;

    width: 100%;

    margin-top: 10px;

}

div.stButton > button:first-child:hover {

    background-color: #E64A19 !important;

}



/* Styling Input Box */

div[data-baseweb="select"] > div, input[type="number"] {

    background-color: #262730 !important;

    color: white !important;

    border-radius: 6px !important;

    border: 1px solid #444 !important;

}

</style>

""", unsafe_allow_html=True)





# --- HELPER FUNCTIONS ---

def get_image_base64(path):

    if os.path.exists(path):

        with open(path, "rb") as image_file:

            return base64.b64encode(image_file.read()).decode()

    return ""



def clean_wisman(teks):
    """Membersihkan angka Wisman, termasuk angka Word yang terpecah baris."""
    if pd.isna(teks):
        return 0
    teks_str = re.sub(r"\s+", "", str(teks)).strip()
    if teks_str in ("", "-"):
        return 0
    try:
        # Format angka Wisman menggunakan titik/koma sebagai pemisah ribuan.
        return int(teks_str.replace(".", "").replace(",", ""))
    except (TypeError, ValueError):
        return 0


def clean_hotel(teks):
    """Membersihkan angka desimal Indonesia dari sel tabel Word."""
    if pd.isna(teks) or str(teks).strip() in ("", "-"):
        return 0.0
    if isinstance(teks, (int, float, np.number)):
        return float(teks)
    teks_str = re.sub(r"\s+", "", str(teks)).strip()
    if "," in teks_str and "." not in teks_str:
        teks_str = teks_str.replace(",", ".")
    elif "," in teks_str and "." in teks_str:
        teks_str = teks_str.replace(".", "").replace(",", ".")
    try:
        return float(teks_str)
    except (TypeError, ValueError):
        return 0.0



def _norm_key(value):
    """Normalisasi label agar pencocokan nama kolom/negara tidak sensitif spasi/tanda baca."""
    return re.sub(r"[^a-z0-9]", "", str(value or "").strip().lower())


def _shift_month(year, month, offset):
    """Geser periode kalender dengan aman melewati pergantian tahun."""
    n = int(year) * 12 + (int(month) - 1) + int(offset)
    return n // 12, (n % 12) + 1


def _find_sheet_for_month(excel_file, prefix, month):
    """Cari sheet bulan berdasarkan nama; mendukung variasi ags/agt dan des/des24."""
    aliases = {
        1: {"jan", "januari"}, 2: {"feb", "februari"},
        3: {"mar", "maret"}, 4: {"apr", "april"},
        5: {"mei", "may"}, 6: {"jun", "juni"},
        7: {"jul", "juli"}, 8: {"ags", "agt", "agu", "aug", "agustus"},
        9: {"sep", "september"}, 10: {"okt", "oct", "oktober", "october"},
        11: {"nov", "november"}, 12: {"des", "dec", "desember", "december"},
    }
    xls = pd.ExcelFile(excel_file)
    prefix_norm = _norm_key(prefix)
    candidates = []
    for sheet in xls.sheet_names:
        normalized = _norm_key(sheet)
        if not normalized.startswith(prefix_norm):
            continue
        suffix = normalized[len(prefix_norm):]
        if suffix in {_norm_key(x) for x in aliases[int(month)]}:
            candidates.append(sheet)
    if len(candidates) != 1:
        raise ValueError(
            f"Sheet untuk '{prefix} {get_short_bulan(int(month))}' tidak ditemukan atau ambigu. "
            f"Kandidat: {candidates or 'tidak ada'}. Periksa nama sheet Excel Wisman."
        )
    return candidates[0]


def _find_header_row_and_column(df, header_label):
    """Cari kolom berdasarkan nama header sebenarnya, bukan indeks kolom tetap."""
    target = _norm_key(header_label)
    for i in range(min(15, len(df))):
        for j, value in enumerate(df.iloc[i].tolist()):
            if _norm_key(value) == target:
                return i, j
    raise ValueError(f"Header '{header_label}' tidak ditemukan pada sheet Excel Wisman.")


def extract_wisman_bulan(file_wisman, bulan):
    """Ambil Wisman Aceh dari kolom SIM dan Sabang yang ditemukan lewat nama header.

    Data negara dicocokkan berdasarkan label kebangsaan. Total dihitung dari seluruh
    baris berkode negara dua huruf pada dua pintu masuk, lalu 'Lainnya' menjadi total
    dikurangi 10 negara utama. Ini menghindari ketergantungan pada nomor baris/kolom tetap.
    """
    sheet_udara = _find_sheet_for_month(file_wisman, "udara", bulan)
    sheet_non_udara = _find_sheet_for_month(file_wisman, "non udara", bulan)
    udara = pd.read_excel(file_wisman, sheet_name=sheet_udara, header=None)
    non_udara = pd.read_excel(file_wisman, sheet_name=sheet_non_udara, header=None)

    header_udara, col_sim = _find_header_row_and_column(udara, "SULTAN ISKANDAR MUDA")
    header_non, col_sabang = _find_header_row_and_column(non_udara, "SABANG")
    # Pada beberapa sheet, label header 'Kebangsaan' tidak ada/bergeser,
    # tetapi struktur data tetap: kode negara di kolom pertama dan nama negara di kolom kedua.
    col_negara_udara = 1
    col_negara_non = 1

    # Label negara bisa berbeda bahasa antara file sumber dan tabel BRS.
    aliases = {
        "malaysia": {"malaysia"},
        "selandia_baru": {"newzealand", "selandiabaru"},
        "perancis": {"france", "perancis"},
        "jerman": {"germany", "jerman"},
        "australia": {"australia"},
        "belanda": {"netherlands", "belanda"},
        "amerika_serikat": {"unitedstatesofamerica", "unitedstates", "amerikaserikat"},
        "singapura": {"singapore", "singapura"},
        "inggris": {"unitedkingdom", "greatbritain", "inggris"},
        "thailand": {"thailand"},
    }

    def rows_by_country(df, header_row, country_col, value_col):
        result = {}
        for ridx in range(header_row + 1, len(df)):
            code = str(df.iloc[ridx, 0]).strip().upper()
            label = _norm_key(df.iloc[ridx, country_col])
            if not re.fullmatch(r"[A-Z]{2}", code):
                continue
            # Skip rows with no usable country label.
            if not label or label in {"nan", "none"}:
                continue
            result[label] = clean_wisman(df.iloc[ridx, value_col])
        return result

    udara_rows = rows_by_country(udara, header_udara, col_negara_udara, col_sim)
    non_rows = rows_by_country(non_udara, header_non, col_negara_non, col_sabang)

    # Pastikan setiap negara utama ada pada kedua sheet, supaya perubahan format tidak
    # diam-diam menghasilkan nol.
    wisman = {}
    for key, labels in aliases.items():
        match_udara = next((label for label in labels if label in udara_rows), None)
        match_non = next((label for label in labels if label in non_rows), None)
        if match_udara is None or match_non is None:
            raise ValueError(
                f"Negara untuk variabel '{key}' tidak ditemukan pada sheet '{sheet_udara}' "
                f"atau '{sheet_non_udara}'. Periksa label kebangsaan sumber."
            )
        wisman[key] = int(udara_rows[match_udara] + non_rows[match_non])

    total_udara = sum(udara_rows.values())
    total_non = sum(non_rows.values())
    total = int(total_udara + total_non)
    jumlah_top10 = sum(wisman.values())
    if jumlah_top10 > total:
        raise ValueError(
            f"Jumlah 10 negara utama ({jumlah_top10}) melebihi total Wisman ({total}). "
            "Periksa kolom sumber yang terpilih."
        )
    wisman["total_wisman"] = total
    wisman["lainnya"] = total - jumlah_top10
    return wisman


def extract_tabel1_generation(doc_lama, tahun_sumber, bulan_sumber):
    """Baca Tabel 1 lama menjadi data bulanan berlabel periode dan kumulatif tahun lalu."""
    monthly_records, warnings = _parse_tabel1_dari_dokumen(
        doc_lama, int(tahun_sumber), int(bulan_sumber)
    )
    if warnings:
        raise ValueError("Tabel 1 BRS sebelumnya tidak valid: " + " | ".join(warnings))

    tabel = doc_lama.tables[0]
    idx_kumulatif = 13 - int(bulan_sumber)
    targets = [
        "Malaysia", "Selandia Baru", "Perancis", "Jerman", "Australia",
        "Belanda", "Amerika Serikat", "Singapura", "Inggris", "Thailand",
        "Lainnya", "JUMLAH",
    ]
    output = {}
    for target in targets:
        row = next(
            (r for r in tabel.rows
             if _norm_key(r.cells[0].text) == _norm_key(target)),
            None,
        )
        if row is None:
            raise ValueError(f"Baris '{target}' tidak ditemukan pada Tabel 1 BRS sebelumnya.")
        numeric_cells = row.cells[-15:]
        monthly = {
            period: int(values[target])
            for period, values in monthly_records.items()
        }
        output[target] = {
            "monthly": monthly,
            "cumulative": clean_wisman(numeric_cells[idx_kumulatif].text),
            "reported_total": clean_wisman(numeric_cells[-1].text),
        }
    return output


def _month_sheet_available(file_wisman, prefix, month):
    try:
        _find_sheet_for_month(file_wisman, prefix, month)
        return True
    except ValueError:
        return False


def reconcile_tabel1_with_raw_wisman(data_tabel1, file_wisman, tahun, bulan_terakhir):
    """Koreksi hanya sel BRS yang terbukti tidak konsisten dengan kolom Total.

    Nilai BRS lama dipertahankan sebagai sumber utama. Raw bulanan dipakai sebagai
    koreksi hanya bila mengganti tepat satu sel akan membuat jumlah Jan s.d. bulan
    sumber sama persis dengan kolom Total. Jika koreksi tidak unik, proses dihentikan
    agar aplikasi tidak menebak-nebak angka.
    """
    keys_by_row = {
        "Malaysia": "malaysia", "Selandia Baru": "selandia_baru", "Perancis": "perancis",
        "Jerman": "jerman", "Australia": "australia", "Belanda": "belanda",
        "Amerika Serikat": "amerika_serikat", "Singapura": "singapura", "Inggris": "inggris",
        "Thailand": "thailand", "Lainnya": "lainnya", "JUMLAH": "total_wisman",
    }
    raw_cache = {}
    for month in range(1, int(bulan_terakhir) + 1):
        has_udara = _month_sheet_available(file_wisman, "udara", month)
        has_non_udara = _month_sheet_available(file_wisman, "non udara", month)
        if not has_udara and not has_non_udara:
            continue
        if has_udara != has_non_udara:
            raise ValueError(
                f"Sheet Wisman {get_short_bulan(month)} {tahun} tidak lengkap: "
                "sheet udara dan non udara harus tersedia berpasangan."
            )
        raw_cache[month] = extract_wisman_bulan(file_wisman, month)

    corrections = []
    periods = [(int(tahun), m) for m in range(1, int(bulan_terakhir) + 1)]
    for nama_baris, key in keys_by_row.items():
        row_data = data_tabel1[nama_baris]
        if not all(period in row_data["monthly"] for period in periods):
            continue
        total_brs = int(row_data.get("reported_total", 0))
        jumlah_brs = sum(int(row_data["monthly"][period]) for period in periods)
        if jumlah_brs == total_brs:
            continue

        selisih = total_brs - jumlah_brs
        kandidat = []
        for month, raw_values in raw_cache.items():
            period = (int(tahun), month)
            nilai_lama = int(row_data["monthly"][period])
            nilai_raw = int(raw_values[key])
            if nilai_raw - nilai_lama == selisih:
                kandidat.append((month, nilai_raw))

        if len(kandidat) == 1:
            month, nilai_raw = kandidat[0]
            nilai_lama = int(row_data["monthly"][(int(tahun), month)])
            row_data["monthly"][(int(tahun), month)] = nilai_raw
            corrections.append(
                f"{nama_baris}: {get_short_bulan(month)} {tahun} dikoreksi "
                f"dari {format_ribuan(nilai_lama)} menjadi {format_ribuan(nilai_raw)} "
                f"agar jumlah Jan–{get_short_bulan(bulan_terakhir)} cocok dengan Total BRS."
            )
        else:
            raise ValueError(
                f"Tabel 1 BRS sebelumnya tidak konsisten pada baris '{nama_baris}': "
                f"jumlah Januari–{get_short_bulan(bulan_terakhir)} {tahun} = "
                f"{format_ribuan(jumlah_brs)}, sedangkan kolom Total BRS = "
                f"{format_ribuan(total_brs)}. Raw workbook tidak memberikan koreksi "
                "satu-sel yang unik. Periksa sumber BRS/raw; aplikasi tidak akan menebak angka."
            )
    return corrections


def validate_tabel1_totals(data_tabel1, tahun_sumber, bulan_sumber):
    """Cegah angka bulanan yang tidak konsisten dengan kolom Total BRS lama."""
    for nama_baris, row_data in data_tabel1.items():
        periode_bulanan = [(int(tahun_sumber), m) for m in range(1, int(bulan_sumber) + 1)]
        # Validasi hanya jika seluruh bulan sampai periode BRS tersedia.
        if not all(p in row_data["monthly"] for p in periode_bulanan):
            continue
        jumlah_bulanan = sum(int(row_data["monthly"][p]) for p in periode_bulanan)
        total_terlapor = int(row_data.get("reported_total", 0))
        if jumlah_bulanan != total_terlapor:
            selisih = jumlah_bulanan - total_terlapor
            raise ValueError(
                f"Tabel 1 BRS sebelumnya tidak konsisten pada baris '{nama_baris}': "
                f"jumlah Januari–{get_short_bulan(bulan_sumber)} {tahun_sumber} setelah dicocokkan "
                f"dengan raw = {format_ribuan(jumlah_bulanan)}, sedangkan kolom Total BRS = "
                f"{format_ribuan(total_terlapor)} (selisih {format_ribuan(selisih)}). "
                "Periksa dokumen sumber atau raw Wisman sebelum Generate agar kesalahan tidak diteruskan."
            )


def build_tabel1_output_row(row_data, nilai_bulan_ini, tahun, bulan, tahun_sumber):
    """Bangun urutan kolom Tabel 1 berdasarkan tahun/bulan, bukan geser indeks mentah."""
    monthly = row_data["monthly"]
    tahun = int(tahun)
    bulan = int(bulan)
    tahun_sumber = int(tahun_sumber)
    values = []

    # Bulan target s.d. Desember tahun sebelumnya.
    for m in range(bulan, 13):
        period = (tahun - 1, m)
        if period not in monthly:
            raise ValueError(
                f"Data Tabel 1 untuk {get_short_bulan(m)} {tahun - 1} tidak tersedia "
                "di BRS sebelumnya. Generate dihentikan agar angka tidak bergeser."
            )
        values.append(int(monthly[period]))

    # Kolom Jan-Des tahun sebelumnya. Untuk BRS Januari, hitung dari 12 bulan
    # tahun sebelumnya yang ada pada BRS Desember; bulan lain gunakan kolom kumulatif.
    if tahun_sumber == tahun:
        kumulatif = int(row_data["cumulative"])
    else:
        periode_tahun_lalu = [(tahun - 1, m) for m in range(1, 13)]
        missing = [p for p in periode_tahun_lalu if p not in monthly]
        if missing:
            raise ValueError(
                "Kolom Jan-Des tahun sebelumnya tidak bisa dihitung karena data bulanan "
                "tahun tersebut tidak lengkap pada BRS sebelumnya."
            )
        kumulatif = sum(int(monthly[p]) for p in periode_tahun_lalu)
    values.append(kumulatif)

    # Januari s.d. bulan berjalan tahun ini.
    current_ytd = 0
    for m in range(1, bulan + 1):
        if m == bulan:
            value = int(nilai_bulan_ini)
        else:
            period = (tahun, m)
            if period not in monthly:
                raise ValueError(
                    f"Data Tabel 1 untuk {get_short_bulan(m)} {tahun} tidak tersedia "
                    "di BRS sebelumnya. Generate dihentikan agar total tidak salah."
                )
            value = int(monthly[period])
        values.append(value)
        current_ytd += value

    values.append(current_ytd)
    if len(values) != 15:
        raise ValueError(f"Jumlah kolom Tabel 1 tidak sesuai: {len(values)} (seharusnya 15).")
    return values


def update_tabel1_year_groups(tabel, tahun, bulan):
    """Sesuaikan merge header tahun Tabel 1 dengan jumlah kolom tiap tahun.

    Struktur numeric: bulan target..Des tahun sebelumnya + Jan-Des tahun sebelumnya,
    lalu Jan..bulan berjalan tahun ini + Total. Jadi span tahun lalu = 14-bulan,
    sedangkan span tahun berjalan = bulan+1.
    """
    row = tabel.rows[0]
    physical_cells = row._tr.tc_lst
    if len(physical_cells) < 3:
        raise ValueError(
            "Struktur header tahun Tabel 1 pada Template.docx tidak sesuai; "
            "diperlukan kolom Kebangsaan dan dua grup tahun."
        )
    span_tahun_lalu = 14 - int(bulan)
    span_tahun_ini = int(bulan) + 1
    if span_tahun_lalu < 1 or span_tahun_ini < 2 or span_tahun_lalu + span_tahun_ini != 15:
        raise ValueError("Jumlah kolom header tahun Tabel 1 tidak valid.")

    for tc, span in ((physical_cells[1], span_tahun_lalu), (physical_cells[2], span_tahun_ini)):
        tc_pr = tc.get_or_add_tcPr()
        grid_span = tc_pr.gridSpan
        if grid_span is None:
            grid_span = OxmlElement("w:gridSpan")
            tc_pr.append(grid_span)
        grid_span.val = span

    def set_text_keep_style(cell, text):
        paragraphs = cell.paragraphs or [cell.add_paragraph()]
        first = paragraphs[0]
        if first.runs:
            first.runs[0].text = str(text)
            for run in first.runs[1:]:
                run.text = ""
        else:
            first.add_run(str(text))
        # Sel yang sebelumnya di-merge dapat menyimpan beberapa paragraf lama.
        # Kosongkan teks sisanya agar label tahun tidak terduplikasi.
        for paragraph in paragraphs[1:]:
            for run in paragraph.runs:
                run.text = ""

    # Setelah span diubah, kolom logis pertama grup tahun ini berada setelah span tahun lalu.
    set_text_keep_style(tabel.cell(0, 1), int(tahun) - 1)
    set_text_keep_style(tabel.cell(0, 1 + span_tahun_lalu), int(tahun))


def period_value(row_data, periode, label):
    """Ambil nilai bulanan dari struktur Tabel 1 dan hentikan proses bila tidak ada."""
    if periode not in row_data["monthly"]:
        raise ValueError(
            f"Data {label} untuk {get_short_bulan(periode[1])} {periode[0]} tidak ditemukan "
            "pada BRS sebelumnya."
        )
    return int(row_data["monthly"][periode])


def format_ribuan(angka): 

    return f"{int(round(float(angka))):,}".replace(",", ".")

def format_desimal(angka, digit=2): 

    return f"{float(angka):,.{digit}f}".replace(",", "X").replace(".", ",").replace("X", ".")

def get_nama_bulan(angka_bulan): 

    return ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober", "November", "Desember"][angka_bulan - 1]

def get_short_bulan(angka_bulan): 

    return ["Jan", "Feb", "Mar", "Apr", "Mei", "Jun", "Jul", "Ags", "Sep", "Okt", "Nov", "Des"][angka_bulan - 1]

def get_arah_tren(angka): 

    return "peningkatan" if angka > 0 else "penurunan"

def get_arah_tren_singkat(angka): 

    return "naik" if angka > 0 else "turun"




# --- GOOGLE SHEETS HISTORIS ---
# Data historis disimpan di Google Sheets agar persisten di Streamlit Cloud.
# Spreadsheet harus dibagikan kepada email service account yang digunakan aplikasi.

HISTORIS_WISMAN_SHEET = "Historis Wisman"
HISTORIS_HOTEL_SHEET = "Historis Hotel"

WISMAN_HEADERS = [
    "tahun", "bulan", "periode", "negara", "nilai", "created_at", "updated_at"
]

HOTEL_HEADERS = [
    "tahun", "bulan", "periode", "kategori", "indikator",
    "hotel_bintang", "akomodasi_lainnya", "created_at", "updated_at"
]


@st.cache_resource(show_spinner=False)
def get_google_sheets_client():
    """Membuat koneksi gspread dari Streamlit Secrets."""
    if "gcp_service_account" not in st.secrets:
        raise RuntimeError(
            "Secrets 'gcp_service_account' belum dikonfigurasi di Streamlit Cloud."
        )

    spreadsheet_id = st.secrets.get("SIBRIO_SPREADSHEET_ID")
    if not spreadsheet_id:
        raise RuntimeError(
            "Secret 'SIBRIO_SPREADSHEET_ID' belum dikonfigurasi di Streamlit Cloud."
        )

    service_account_info = dict(st.secrets["gcp_service_account"])
    scopes = [
        "https://www.googleapis.com/auth/spreadsheets",
        "https://www.googleapis.com/auth/drive",
    ]
    credentials = Credentials.from_service_account_info(
        service_account_info,
        scopes=scopes,
    )
    return gspread.authorize(credentials)


@st.cache_resource(show_spinner=False)
def get_historis_spreadsheet():
    """Membuka spreadsheet historis berdasarkan Spreadsheet ID."""
    client = get_google_sheets_client()
    spreadsheet_id = st.secrets["SIBRIO_SPREADSHEET_ID"]
    return client.open_by_key(spreadsheet_id)


def get_or_create_worksheet(title, headers):
    """Mengambil worksheet historis atau membuatnya jika belum ada."""
    spreadsheet = get_historis_spreadsheet()
    try:
        worksheet = spreadsheet.worksheet(title)
    except gspread.WorksheetNotFound:
        worksheet = spreadsheet.add_worksheet(title=title, rows=1000, cols=max(10, len(headers)))
        worksheet.append_row(headers, value_input_option="USER_ENTERED")
        return worksheet

    values = worksheet.get_all_values()
    if not values:
        worksheet.append_row(headers, value_input_option="USER_ENTERED")
    elif values[0] != headers:
        # Jangan menimpa data lama. Jika header berbeda, hentikan agar struktur data tidak rusak.
        raise RuntimeError(
            f"Header worksheet '{title}' tidak sesuai dengan struktur SIBRIO. "
            "Periksa worksheet tersebut sebelum melanjutkan."
        )
    return worksheet


def _find_matching_row(values, key_indexes, key_values):
    """Mencari nomor baris Google Sheets berdasarkan beberapa kolom kunci."""
    if not values:
        return None
    for row_number, row in enumerate(values[1:], start=2):
        # Pad row agar aman jika ada sel kosong di ujung.
        row = list(row) + [""] * (max(key_indexes) + 1 - len(row))
        if all(str(row[idx]).strip() == str(value).strip() for idx, value in zip(key_indexes, key_values)):
            return row_number
    return None


def _upsert_rows(worksheet, records, key_indexes):
    """UPSERT sederhana ke Google Sheets: update baris jika ada, append jika belum."""
    values = worksheet.get_all_values()
    updates = []
    appends = []

    for record in records:
        row_number = _find_matching_row(
            values,
            key_indexes,
            [record[idx] for idx in key_indexes],
        )
        if row_number is None:
            appends.append(record)
        else:
            updates.append((row_number, record))

    # Update record yang sudah ada.
    for row_number, record in updates:
        end_col = len(record)
        worksheet.update(
            f"A{row_number}:{gspread.utils.rowcol_to_a1(row_number, end_col).rstrip('0123456789')}{row_number}",
            [record],
            value_input_option="USER_ENTERED",
        )

    # Tambahkan record baru sekaligus.
    if appends:
        worksheet.append_rows(appends, value_input_option="USER_ENTERED")


def save_historis_wisman(tahun, bulan, wisman_data):
    """UPSERT data Wisman bulan berjalan ke Google Sheets tanpa menghapus bulan lain."""
    negara_map = [
        ("Malaysia", "malaysia"),
        ("Selandia Baru", "selandia_baru"),
        ("Perancis", "perancis"),
        ("Jerman", "jerman"),
        ("Australia", "australia"),
        ("Belanda", "belanda"),
        ("Amerika Serikat", "amerika_serikat"),
        ("Singapura", "singapura"),
        ("Inggris", "inggris"),
        ("Thailand", "thailand"),
        ("Lainnya", "lainnya"),
        ("JUMLAH", "total_wisman"),
    ]

    worksheet = get_or_create_worksheet(HISTORIS_WISMAN_SHEET, WISMAN_HEADERS)
    sekarang = datetime.now().isoformat(timespec="seconds")
    periode = f"{get_nama_bulan(int(bulan))} {int(tahun)}"

    records = []
    for negara, key in negara_map:
        nilai = int(round(float(wisman_data.get(key, 0) or 0)))
        records.append([
            int(tahun), int(bulan), periode, negara, nilai, sekarang, sekarang
        ])

    _upsert_rows(worksheet, records, key_indexes=[0, 1, 3])



def get_rolling_month_periods(brs_tahun, brs_bulan):
    """Menghasilkan 13 periode bulanan yang benar-benar tersimpan di Tabel 1 BRS.

    Struktur Tabel 1 SIBRIO:
    [bulan akhir tahun sebelumnya ... Des tahun sebelumnya,
     Jan ... bulan berjalan tahun berjalan]
    dengan satu kolom Jan-Des dan satu kolom Total yang tidak dihitung sebagai
    observasi bulanan.
    """
    periods = []
    # Bagian sebelum Jan-Des: bulan berjalan s.d. Desember tahun sebelumnya.
    for month in range(int(brs_bulan), 13):
        periods.append((int(brs_tahun) - 1, month))
    # Bagian setelah Jan-Des: Januari s.d. bulan BRS tahun berjalan.
    for month in range(1, int(brs_bulan) + 1):
        periods.append((int(brs_tahun), month))
    return periods


def backfill_historis_tabel1_dari_brs(doc_lama, brs_tahun, brs_bulan):
    """Mengimpor seluruh observasi bulanan Tabel 1 yang tersedia di BRS lama.

    Contoh BRS Juni 2026 akan mengisi data Tabel 1:
    Juni-Desember 2025 + Januari-Juni 2026.
    BRS Desember 2025 akan mengisi Desember 2024 + Januari-Desember 2025.
    Data yang sudah ada di Google Sheets tidak diduplikasi karena UPSERT.
    """
    tabel1 = doc_lama.tables[0]
    if len(tabel1.rows) < 2:
        return 0

    # Header dan data memakai 15 cell terakhir, sama seperti logika existing.
    header_cells = tabel1.rows[1].cells[-15:]
    headers = [str(cell.text).strip() for cell in header_cells]

    idx_jandes = 13 - int(brs_bulan)
    monthly_indices = [i for i in range(14) if i != idx_jandes]
    periods = get_rolling_month_periods(brs_tahun, brs_bulan)

    if len(monthly_indices) != len(periods):
        return 0

    target_negara = [
        "Malaysia", "Selandia Baru", "Perancis", "Jerman", "Australia",
        "Belanda", "Amerika Serikat", "Singapura", "Inggris", "Thailand",
        "Lainnya", "JUMLAH"
    ]

    # Ambil nilai per negara dari BRS lama.
    extracted = {}
    for row in tabel1.rows:
        teks_baris = " ".join(str(c.text) for c in row.cells).upper()
        for target in target_negara:
            if target.upper() in teks_baris and target not in extracted:
                cells = row.cells[-15:]
                values = [clean_wisman(cell.text) for cell in cells]
                extracted[target] = values
                break

    if not extracted:
        return 0

    # Konversi ke struktur wisman_data per periode agar fungsi save existing
    # tetap menjadi satu-satunya pintu UPSERT ke Google Sheets.
    total_saved = 0
    for period_index, (tahun, bulan) in enumerate(periods):
        cell_index = monthly_indices[period_index]
        wisman_data = {}
        key_map = {
            "Malaysia": "malaysia",
            "Selandia Baru": "selandia_baru",
            "Perancis": "perancis",
            "Jerman": "jerman",
            "Australia": "australia",
            "Belanda": "belanda",
            "Amerika Serikat": "amerika_serikat",
            "Singapura": "singapura",
            "Inggris": "inggris",
            "Thailand": "thailand",
            "Lainnya": "lainnya",
            "JUMLAH": "total_wisman",
        }
        for negara, values in extracted.items():
            if negara in key_map and cell_index < len(values):
                wisman_data[key_map[negara]] = values[cell_index]

        if wisman_data:
            save_historis_wisman(tahun, bulan, wisman_data)
            total_saved += 1

    return total_saved


def backfill_historis_tabel2_dari_brs(doc_lama, brs_tahun, brs_bulan, data_word):
    """Mengimpor periode Tabel 2 yang memang tersedia pada BRS lama.

    Struktur existing SIBRIO menggunakan tiga periode pada Tabel 2:
    dua bulan sebelumnya + bulan berjalan. Karena itu fungsi ini tidak
    mengarang data bulan yang tidak terdapat di dokumen.
    """
    # Dokumen BRS lama menampilkan tiga bulan: bulan-2, bulan-1, dan bulan.
    # data_word['baris_lama_1'] mengambil baris kedua (bulan-1), sedangkan
    # data_word['baris_lama_2'] mengambil baris ketiga (bulan dokumen lama).
    # Karena itu kedua baris tersebut harus disimpan pada bulan-1 dan bulan.
    periode1_num = int(brs_tahun) * 12 + int(brs_bulan) - 1
    periode2_num = int(brs_tahun) * 12 + int(brs_bulan)
    periode3_num = int(brs_tahun) * 12 + int(brs_bulan) + 1

    def from_num(n):
        # Representasi bulan dengan basis 1.
        tahun = (n - 1) // 12
        bulan = ((n - 1) % 12) + 1
        return tahun, bulan

    p1 = from_num(periode1_num)
    p2 = from_num(periode2_num)
    p3 = from_num(periode3_num)

    # Nilai periode 3 diambil dari data Excel bulan berjalan BRS yang sedang
    # dibuat. Saat fungsi ini dipanggil sebelum save bulan berjalan, p3 belum
    # perlu disimpan di sini; fungsi utama akan melakukan save bulan berjalan.
    # Backfill bulan-1 dan bulan dari dokumen lama; periode berikutnya
    # (bulan yang sedang dibuat) akan disimpan dari Excel bulan berjalan.
    periode_lama = [
        (p1, "baris_lama_1"),
        (p2, "baris_lama_2"),
    ]

    jumlah_disimpan = 0
    for (tahun, bulan), nama_baris in periode_lama:
        hotel_data_lama = {
            "rlm_asing": data_word["tabel2"]["asing"][nama_baris][0],
            "rlm_asing_non": data_word["tabel2"]["asing"][nama_baris][1],
            "rlm_nus": data_word["tabel2"]["nusantara"][nama_baris][0],
            "rlm_nus_non": data_word["tabel2"]["nusantara"][nama_baris][1],
            "rlm_total": data_word["tabel2"]["total"][nama_baris][0],
            "rlm_total_non": data_word["tabel2"]["total"][nama_baris][1],
            "tpk_bintang": data_word["tabel2"]["tpk"][nama_baris][0],
            "tpk_non": data_word["tabel2"]["tpk"][nama_baris][1],
        }
        save_historis_hotel(tahun, bulan, hotel_data_lama)
        jumlah_disimpan += 1

    return jumlah_disimpan


def backfill_historis_dari_brs(doc_lama, brs_tahun, brs_bulan, data_word):
    """Backfill histori dari BRS lama tanpa mengganggu data BRS berjalan."""
    hasil = {"tabel1": 0, "tabel2": 0}
    hasil["tabel1"] = backfill_historis_tabel1_dari_brs(
        doc_lama, brs_tahun, brs_bulan
    )
    hasil["tabel2"] = backfill_historis_tabel2_dari_brs(
        doc_lama, brs_tahun, brs_bulan, data_word
    )
    return hasil


def save_historis_hotel(tahun, bulan, hotel_data):
    """UPSERT data Hotel bulan berjalan ke Google Sheets tanpa menghapus bulan lain."""
    rows = [
        ("Rata-Rata Lama Menginap", "Asing",
         hotel_data.get("rlm_asing", 0), hotel_data.get("rlm_asing_non", 0)),
        ("Rata-Rata Lama Menginap", "Nusantara",
         hotel_data.get("rlm_nus", 0), hotel_data.get("rlm_nus_non", 0)),
        ("Rata-Rata Lama Menginap", "Total",
         hotel_data.get("rlm_total", 0), hotel_data.get("rlm_total_non", 0)),
        ("TPK", "TPK",
         hotel_data.get("tpk_bintang", 0), hotel_data.get("tpk_non", 0)),
    ]

    worksheet = get_or_create_worksheet(HISTORIS_HOTEL_SHEET, HOTEL_HEADERS)
    sekarang = datetime.now().isoformat(timespec="seconds")
    periode = f"{get_nama_bulan(int(bulan))} {int(tahun)}"

    records = []
    for kategori, indikator, hotel_bintang, akomodasi_lainnya in rows:
        records.append([
            int(tahun), int(bulan), periode, kategori, indikator,
            float(hotel_bintang or 0), float(akomodasi_lainnya or 0),
            sekarang, sekarang,
        ])

    _upsert_rows(worksheet, records, key_indexes=[0, 1, 3, 4])


def _worksheet_to_dataframe(title, headers):
    worksheet = get_or_create_worksheet(title, headers)
    values = worksheet.get_all_records()
    return pd.DataFrame(values, columns=headers) if values else pd.DataFrame(columns=headers)


def get_historis_wisman(tahun_awal, bulan_awal, tahun_akhir, bulan_akhir):
    """Mengambil historis Wisman dari Google Sheets sesuai rentang periode."""
    df = _worksheet_to_dataframe(HISTORIS_WISMAN_SHEET, WISMAN_HEADERS)
    if df.empty:
        return df

    df["tahun"] = pd.to_numeric(df["tahun"], errors="coerce")
    df["bulan"] = pd.to_numeric(df["bulan"], errors="coerce")
    df["nilai"] = pd.to_numeric(df["nilai"], errors="coerce").fillna(0).astype(int)
    df = df.dropna(subset=["tahun", "bulan"])

    periode_num = df["tahun"] * 12 + df["bulan"]
    awal = int(tahun_awal) * 12 + int(bulan_awal)
    akhir = int(tahun_akhir) * 12 + int(bulan_akhir)
    df = df[(periode_num >= awal) & (periode_num <= akhir)].copy()

    urutan_negara = [
        "Malaysia", "Selandia Baru", "Perancis", "Jerman", "Australia",
        "Belanda", "Amerika Serikat", "Singapura", "Inggris", "Thailand",
        "Lainnya", "JUMLAH"
    ]
    urutan = {nama: i for i, nama in enumerate(urutan_negara)}
    df["_urutan"] = df["negara"].map(urutan).fillna(999)
    return df.sort_values(["tahun", "bulan", "_urutan"]).drop(columns=["_urutan"])


def get_historis_hotel(tahun_awal, bulan_awal, tahun_akhir, bulan_akhir):
    """Mengambil historis Hotel dari Google Sheets sesuai rentang periode."""
    df = _worksheet_to_dataframe(HISTORIS_HOTEL_SHEET, HOTEL_HEADERS)
    if df.empty:
        return df

    for col in ["tahun", "bulan"]:
        df[col] = pd.to_numeric(df[col], errors="coerce")
    for col in ["hotel_bintang", "akomodasi_lainnya"]:
        df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0.0)
    df = df.dropna(subset=["tahun", "bulan"])

    periode_num = df["tahun"] * 12 + df["bulan"]
    awal = int(tahun_awal) * 12 + int(bulan_awal)
    akhir = int(tahun_akhir) * 12 + int(bulan_akhir)
    df = df[(periode_num >= awal) & (periode_num <= akhir)].copy()

    urutan_kategori = {"Rata-Rata Lama Menginap": 1, "TPK": 2}
    urutan_indikator = {"Asing": 1, "Nusantara": 2, "Total": 3, "TPK": 4}
    df["_kategori"] = df["kategori"].map(urutan_kategori).fillna(999)
    df["_indikator"] = df["indikator"].map(urutan_indikator).fillna(999)
    return df.sort_values(["tahun", "bulan", "_kategori", "_indikator"]).drop(
        columns=["_kategori", "_indikator"]
    )


def buat_excel_historis(df_wisman, df_hotel, pilihan):
    """Membuat file Excel historis sesuai pilihan user."""
    output = io.BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        sheet_dibuat = False

        if pilihan in ("Tabel 1 - Wisman", "Tabel 1 + Tabel 2") and not df_wisman.empty:
            pivot_wisman = df_wisman.pivot(
                index="negara", columns="periode", values="nilai"
            )
            periode_order = (
                df_wisman[["tahun", "bulan", "periode"]]
                .drop_duplicates()
                .sort_values(["tahun", "bulan"])["periode"]
                .tolist()
            )
            urutan_negara = [
                "Malaysia", "Selandia Baru", "Perancis", "Jerman",
                "Australia", "Belanda", "Amerika Serikat", "Singapura",
                "Inggris", "Thailand", "Lainnya", "JUMLAH"
            ]
            pivot_wisman = pivot_wisman.reindex(
                index=[x for x in urutan_negara if x in pivot_wisman.index],
                columns=periode_order
            )
            pivot_wisman = pivot_wisman.fillna("")
            pivot_wisman.to_excel(writer, sheet_name="Historis Wisman")
            sheet_dibuat = True

        if pilihan in ("Tabel 2 - Hotel", "Tabel 1 + Tabel 2") and not df_hotel.empty:
            df_h = df_hotel.copy()
            periode_order = (
                df_h[["tahun", "bulan", "periode"]]
                .drop_duplicates()
                .sort_values(["tahun", "bulan"])["periode"]
                .tolist()
            )

            rows = []
            for _, r in df_h.iterrows():
                rows.append({
                    "Kategori": r["kategori"],
                    "Indikator": r["indikator"],
                    "Jenis": "Hotel Bintang",
                    "Nilai": r["hotel_bintang"],
                    "tahun": r["tahun"],
                    "bulan": r["bulan"],
                    "periode": r["periode"],
                })
                rows.append({
                    "Kategori": r["kategori"],
                    "Indikator": r["indikator"],
                    "Jenis": "Akomodasi Lainnya",
                    "Nilai": r["akomodasi_lainnya"],
                    "tahun": r["tahun"],
                    "bulan": r["bulan"],
                    "periode": r["periode"],
                })

            long_hotel = pd.DataFrame(rows)
            pivot_hotel = long_hotel.pivot_table(
                index=["Kategori", "Indikator", "Jenis"],
                columns="periode",
                values="Nilai",
                aggfunc="first"
            )
            pivot_hotel = pivot_hotel.reindex(columns=periode_order)
            pivot_hotel = pivot_hotel.reset_index()
            pivot_hotel.to_excel(writer, sheet_name="Historis Hotel", index=False)
            sheet_dibuat = True

        if not sheet_dibuat:
            pd.DataFrame({
                "Informasi": ["Tidak ada data historis untuk periode yang dipilih."]
            }).to_excel(writer, sheet_name="Informasi", index=False)

    output.seek(0)
    return output


def tampilkan_historis_data():
    """UI untuk melihat dan mengunduh data historis dari Google Sheets."""
    st.divider()
    st.markdown("## 📊 Historis Data")
    st.caption(
        "Data historis SIBRIO disimpan di Google Sheets sehingga tetap tersedia "
        "meskipun aplikasi Streamlit Cloud restart atau redeploy."
    )

    try:
        # Uji koneksi agar error konfigurasi tampil jelas di bagian historis.
        get_historis_spreadsheet()
    except Exception as e:
        st.warning(
            "⚠️ Fitur Historis Data belum terhubung ke Google Sheets. "
            f"Periksa Streamlit Secrets dan akses spreadsheet. Detail: {e}"
        )
        return

    pilihan = st.selectbox(
        "Pilih data yang ingin dilihat/download",
        ["Tabel 1 - Wisman", "Tabel 2 - Hotel", "Tabel 1 + Tabel 2"],
        key="historis_pilihan"
    )

    bulan_nama = [
        "Januari", "Februari", "Maret", "April", "Mei", "Juni",
        "Juli", "Agustus", "September", "Oktober", "November", "Desember"
    ]

    col1, col2 = st.columns(2)
    with col1:
        st.markdown("**Periode Awal**")
        c1, c2 = st.columns(2)
        with c1:
            bulan_awal = st.selectbox(
                "Bulan awal", list(range(1, 13)),
                format_func=lambda x: bulan_nama[x - 1],
                index=0, key="hist_bulan_awal"
            )
        with c2:
            tahun_awal = st.number_input(
                "Tahun awal", min_value=2020, max_value=2050,
                value=2025, step=1, key="hist_tahun_awal"
            )

    with col2:
        st.markdown("**Periode Akhir**")
        c3, c4 = st.columns(2)
        with c3:
            bulan_akhir = st.selectbox(
                "Bulan akhir", list(range(1, 13)),
                format_func=lambda x: bulan_nama[x - 1],
                index=4, key="hist_bulan_akhir"
            )
        with c4:
            tahun_akhir = st.number_input(
                "Tahun akhir", min_value=2020, max_value=2050,
                value=2026, step=1, key="hist_tahun_akhir"
            )

    awal = int(tahun_awal) * 12 + int(bulan_awal)
    akhir = int(tahun_akhir) * 12 + int(bulan_akhir)

    if awal > akhir:
        st.error("❌ Periode awal tidak boleh lebih besar daripada periode akhir.")
        return

    try:
        df_wisman = pd.DataFrame()
        df_hotel = pd.DataFrame()

        if pilihan in ("Tabel 1 - Wisman", "Tabel 1 + Tabel 2"):
            df_wisman = get_historis_wisman(
                tahun_awal, bulan_awal, tahun_akhir, bulan_akhir
            )

        if pilihan in ("Tabel 2 - Hotel", "Tabel 1 + Tabel 2"):
            df_hotel = get_historis_hotel(
                tahun_awal, bulan_awal, tahun_akhir, bulan_akhir
            )
    except Exception as e:
        st.error(f"❌ Gagal membaca data historis dari Google Sheets: {e}")
        return

    ada_data = not df_wisman.empty or not df_hotel.empty

    if not ada_data:
        st.info(
            "ℹ️ Belum terdapat data historis untuk periode yang dipilih. "
            "Data akan otomatis tersimpan setiap kali Generate BRS berhasil dilakukan."
        )
        return

    if not df_wisman.empty:
        st.markdown("### Tabel 1 — Historis Wisman")
        pivot_wisman = df_wisman.pivot(
            index="negara", columns="periode", values="nilai"
        )
        periode_order = (
            df_wisman[["tahun", "bulan", "periode"]]
            .drop_duplicates()
            .sort_values(["tahun", "bulan"])["periode"]
            .tolist()
        )
        urutan_negara = [
            "Malaysia", "Selandia Baru", "Perancis", "Jerman",
            "Australia", "Belanda", "Amerika Serikat", "Singapura",
            "Inggris", "Thailand", "Lainnya", "JUMLAH"
        ]
        pivot_wisman = pivot_wisman.reindex(
            index=[x for x in urutan_negara if x in pivot_wisman.index],
            columns=periode_order
        )
        st.dataframe(pivot_wisman.fillna(""), use_container_width=True)

    if not df_hotel.empty:
        st.markdown("### Tabel 2 — Historis Hotel")
        df_h = df_hotel.copy()
        periode_order = (
            df_h[["tahun", "bulan", "periode"]]
            .drop_duplicates()
            .sort_values(["tahun", "bulan"])["periode"]
            .tolist()
        )
        preview_rows = []
        for _, r in df_h.iterrows():
            preview_rows.append({
                "Kategori": r["kategori"],
                "Indikator": r["indikator"],
                "Jenis": "Hotel Bintang",
                "periode": r["periode"],
                "nilai": r["hotel_bintang"],
            })
            preview_rows.append({
                "Kategori": r["kategori"],
                "Indikator": r["indikator"],
                "Jenis": "Akomodasi Lainnya",
                "periode": r["periode"],
                "nilai": r["akomodasi_lainnya"],
            })

        preview_df = pd.DataFrame(preview_rows)
        preview_pivot = preview_df.pivot_table(
            index=["Kategori", "Indikator", "Jenis"],
            columns="periode",
            values="nilai",
            aggfunc="first"
        ).reindex(columns=periode_order).reset_index()

        st.dataframe(preview_pivot.fillna(""), use_container_width=True)

    nama_awal = f"{get_short_bulan(int(bulan_awal))}{int(tahun_awal)}"
    nama_akhir = f"{get_short_bulan(int(bulan_akhir))}{int(tahun_akhir)}"
    nama_download = f"Historis_SIBRIO_{nama_awal}_{nama_akhir}.xlsx"

    excel_data = buat_excel_historis(df_wisman, df_hotel, pilihan)

    st.download_button(
        "📥 Download Historis Excel",
        data=excel_data,
        file_name=nama_download,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True
    )


# --- MIGRASI HISTORIS DARI DOKUMEN BRS LAMA ---

BULAN_ALIAS = {
    "jan": 1, "januari": 1,
    "feb": 2, "februari": 2,
    "mar": 3, "maret": 3,
    "apr": 4, "april": 4,
    "mei": 5, "may": 5,
    "jun": 6, "juni": 6,
    "jul": 7, "juli": 7,
    "agu": 8, "ags": 8, "agustus": 8, "aug": 8, "august": 8,
    "sep": 9, "september": 9,
    "okt": 10, "oct": 10, "oktober": 10, "october": 10,
    "nov": 11, "november": 11,
    "des": 12, "dec": 12, "desember": 12, "december": 12,
}


def parse_bulan_tahun(teks):
    """Membaca label periode seperti 'Mei 2026', 'Mar'26', atau 'Mar-26'."""
    teks = str(teks or "").strip().lower().replace("’", "'")
    pola = r"(?<![a-z])([a-z]{3,9})\s*[-/' ]\s*'?\s*(20\d{2}|\d{2})(?!\d)"
    for match in re.finditer(pola, teks):
        nama, tahun_teks = match.groups()
        bulan = BULAN_ALIAS.get(nama)
        if bulan:
            tahun = int(tahun_teks)
            if tahun < 100:
                tahun += 2000 if tahun < 70 else 1900
            return tahun, bulan
    # Label tanpa pemisah, misalnya Mei'26
    pola2 = r"(?<![a-z])([a-z]{3,9})\s*'?\s*(20\d{2}|\d{2})(?!\d)"
    for match in re.finditer(pola2, teks):
        nama, tahun_teks = match.groups()
        bulan = BULAN_ALIAS.get(nama)
        if bulan:
            tahun = int(tahun_teks)
            if tahun < 100:
                tahun += 2000 if tahun < 70 else 1900
            return tahun, bulan
    return None


def tebak_periode_brs(file_brs, doc):
    """Menentukan periode dokumen dari nama file, lalu teks dokumen sebagai fallback."""
    nama_file = getattr(file_brs, "name", "") or ""
    # Nama keluaran SIBRIO: BRS_Pariwisata_{bulan}_{tahun}.docx
    match = re.search(r"(?:pariwisata[_\s-]*)?(\d{1,2})[_\s-]+(20\d{2})", nama_file, re.I)
    if match:
        bulan, tahun = int(match.group(1)), int(match.group(2))
        if 1 <= bulan <= 12:
            return tahun, bulan

    potongan_teks = []
    potongan_teks.extend(p.text for p in doc.paragraphs)
    for table in doc.tables[:2]:
        for row in table.rows[:4]:
            potongan_teks.extend(cell.text for cell in row.cells)
    teks = " ".join(potongan_teks)
    # Ambil pasangan bulan-tahun pertama yang masuk akal dari judul/header.
    for match in re.finditer(r"([A-Za-z]{3,9})\s*[,/' -]?\s*(20\d{2})", teks):
        bulan = BULAN_ALIAS.get(match.group(1).lower())
        if bulan:
            return int(match.group(2)), bulan
    return None


def _parse_tabel1_dari_dokumen(doc, tahun_brs, bulan_brs):
    """Mengembalikan observasi bulanan Tabel 1; kolom Jan-Des/Total diabaikan."""
    if not doc.tables:
        return {}, ["Tabel 1 tidak ditemukan"]
    tabel = doc.tables[0]
    if len(tabel.rows) < 2:
        return {}, ["Struktur Tabel 1 tidak dikenali"]

    headers = [str(c.text).strip() for c in tabel.rows[1].cells[-15:]]
    if len(headers) != 15:
        return {}, [f"Tabel 1 memiliki {len(headers)} kolom terakhir, seharusnya 15"]

    periods = get_rolling_month_periods(tahun_brs, bulan_brs)
    idx_jandes = 13 - int(bulan_brs)
    monthly_indices = [i for i in range(14) if i != idx_jandes]
    if len(periods) != len(monthly_indices):
        return {}, ["Header periode Tabel 1 tidak cocok dengan periode BRS"]

    # Validasi label bulan pada header untuk mencegah salah pemetaan lintas tahun.
    for period, cell_idx in zip(periods, monthly_indices):
        label = re.sub(r"[^a-z]", "", headers[cell_idx].lower())
        bulan_header = BULAN_ALIAS.get(label)
        if bulan_header is None:
            return {}, [f"Label bulan Tabel 1 tidak dikenali: '{headers[cell_idx]}'"]
        if bulan_header != period[1]:
            return {}, [
                f"Header Tabel 1 tidak cocok: kolom '{headers[cell_idx]}' "
                f"tidak sesuai periode {get_short_bulan(period[1])} {period[0]}. "
                "Data periode ini tidak diimpor agar tidak salah bulan."
            ]

    targets = [
        "Malaysia", "Selandia Baru", "Perancis", "Jerman", "Australia",
        "Belanda", "Amerika Serikat", "Singapura", "Inggris", "Thailand",
        "Lainnya", "JUMLAH",
    ]
    extracted = {}
    for row in tabel.rows:
        # Nama negara berada di kolom sebelah kiri, di luar 15 kolom angka.
        row_texts = [re.sub(r"\s+", " ", str(c.text)).strip().casefold() for c in row.cells]
        for target in targets:
            if target.casefold() in row_texts and target not in extracted:
                cells = row.cells[-15:]
                extracted[target] = [clean_wisman(c.text) for c in cells]
                break

    missing = [target for target in targets if target not in extracted]
    if missing:
        return {}, ["Baris Tabel 1 tidak ditemukan: " + ", ".join(missing)]

    records = {}
    for j, period in enumerate(periods):
        cell_idx = monthly_indices[j]
        records[period] = {target: extracted[target][cell_idx] for target in targets}
    return records, []


def _parse_tabel2_dari_dokumen(doc, periode_brs=None):
    """Ekstrak semua 3 periode x 4 indikator dari Tabel 2 berdasarkan label baris."""
    if len(doc.tables) < 2:
        return {}, ["Tabel 2 tidak ditemukan"]
    tabel = doc.tables[1]
    if len(tabel.rows) < 12:
        return {}, ["Tabel 2 tidak memiliki 12 baris data yang diharapkan"]

    data_rows = tabel.rows[-12:]
    indikator = ["asing", "nusantara", "total", "tpk"]
    keys = {
        "asing": ("rlm_asing", "rlm_asing_non"),
        "nusantara": ("rlm_nus", "rlm_nus_non"),
        "total": ("rlm_total", "rlm_total_non"),
        "tpk": ("tpk_bintang", "tpk_non"),
    }
    periods_found = {}
    warnings = []

    # Jika label bulan pada sel tidak terbaca, urutan tiga baris ditambatkan
    # ke periode BRS (dua bulan sebelumnya dan bulan BRS itu sendiri).
    fallback_periods = []
    if periode_brs:
        y, m = periode_brs
        for offset in (-2, -1, 0):
            n = y * 12 + m + offset
            fallback_periods.append(((n - 1) // 12, ((n - 1) % 12) + 1))

    for group_idx, nama_indikator in enumerate(indikator):
        for j in range(3):
            row = data_rows[group_idx * 3 + j]
            cells = row.cells
            if len(cells) < 3:
                warnings.append(f"Baris Tabel 2 ke-{group_idx * 3 + j + 1} tidak lengkap")
                continue
            periode = parse_bulan_tahun(cells[-3].text)
            if not periode and len(fallback_periods) == 3:
                periode = fallback_periods[j]
            if not periode:
                warnings.append("Label bulan/tahun Tabel 2 tidak dapat dikenali")
                continue
            nilai_bintang_text = str(cells[-2].text).strip()
            nilai_non_text = str(cells[-1].text).strip()
            if not nilai_bintang_text or not nilai_non_text:
                warnings.append(f"Nilai Tabel 2 kosong untuk {periode[1]}/{periode[0]}")
                continue
            item = periods_found.setdefault(periode, {})
            k1, k2 = keys[nama_indikator]
            item[k1] = clean_hotel(nilai_bintang_text)
            item[k2] = clean_hotel(nilai_non_text)

    required = set(sum((list(v) for v in keys.values()), []))
    hasil = {periode: nilai for periode, nilai in periods_found.items() if required.issubset(nilai)}
    for periode, nilai in periods_found.items():
        if not required.issubset(nilai):
            warnings.append(f"Tabel 2 periode {periode[1]}/{periode[0]} tidak lengkap; tidak disimpan")
    return hasil, warnings


def migrasikan_dokumen_brs(file_brs):
    """Impor satu BRS lama ke historis. Upsert membuat impor aman diulang."""
    doc = docx.Document(file_brs)
    periode_brs = tebak_periode_brs(file_brs, doc)
    if not periode_brs:
        return {"file": getattr(file_brs, "name", "BRS"), "wisman": 0, "hotel": 0,
                "warnings": ["Periode BRS tidak terdeteksi. Gunakan nama file seperti BRS_Pariwisata_4_2026.docx."]}

    hasil = {"file": getattr(file_brs, "name", "BRS"), "wisman": 0, "hotel": 0, "warnings": []}
    wisman_periods, warn_wisman = _parse_tabel1_dari_dokumen(doc, *periode_brs)
    hasil["warnings"].extend(["Tabel 1: " + w for w in warn_wisman])
    for (tahun, bulan), nilai_per_negara in wisman_periods.items():
        # Simpan hanya periode yang berisi seluruh baris target; nol hanya diterima
        # jika memang tertulis pada dokumen sumber.
        data = {
            "malaysia": nilai_per_negara["Malaysia"],
            "selandia_baru": nilai_per_negara["Selandia Baru"],
            "perancis": nilai_per_negara["Perancis"],
            "jerman": nilai_per_negara["Jerman"],
            "australia": nilai_per_negara["Australia"],
            "belanda": nilai_per_negara["Belanda"],
            "amerika_serikat": nilai_per_negara["Amerika Serikat"],
            "singapura": nilai_per_negara["Singapura"],
            "inggris": nilai_per_negara["Inggris"],
            "thailand": nilai_per_negara["Thailand"],
            "lainnya": nilai_per_negara["Lainnya"],
            "total_wisman": nilai_per_negara["JUMLAH"],
        }
        save_historis_wisman(tahun, bulan, data)
        hasil["wisman"] += 1

    hotel_periods, warn_hotel = _parse_tabel2_dari_dokumen(doc, periode_brs)
    hasil["warnings"].extend(["Tabel 2: " + w for w in warn_hotel])
    for (tahun, bulan), hotel_data in hotel_periods.items():
        save_historis_hotel(tahun, bulan, hotel_data)
        hasil["hotel"] += 1
    return hasil


def tampilkan_migrasi_historis():
    """UI migrasi manual; tidak berjalan otomatis saat aplikasi dibuka."""
    st.divider()
    st.markdown("## 🗂️ Migrasi Historis dari BRS Lama")
    st.caption(
        "Unggah satu atau beberapa BRS Word lama untuk memulihkan angka historis. "
        "Impor dapat diulang karena data periode yang sama diperbarui, bukan digandakan. "
        "Untuk Tabel 1, nama file disarankan mengikuti format BRS_Pariwisata_4_2026.docx."
    )
    files_brs = st.file_uploader(
        "Pilih dokumen BRS lama (.docx)", type=["docx"], accept_multiple_files=True,
        key="migrasi_brs_historis_files"
    )
    if st.button("📥 Impor / Pulihkan Historis dari BRS", key="btn_migrasi_historis"):
        if not files_brs:
            st.warning("Pilih minimal satu dokumen BRS terlebih dahulu.")
            return
        try:
            get_historis_spreadsheet()
        except Exception as e:
            st.error(f"Google Sheets belum dapat diakses. Periksa Secrets dan izin spreadsheet. Detail: {e}")
            return
        total_wisman = 0
        total_hotel = 0
        for file_brs in files_brs:
            try:
                file_brs.seek(0)
                hasil = migrasikan_dokumen_brs(file_brs)
                total_wisman += hasil["wisman"]
                total_hotel += hasil["hotel"]
                st.write(
                    f"**{hasil['file']}** — Tabel 1: {hasil['wisman']} periode; "
                    f"Tabel 2: {hasil['hotel']} periode."
                )
                for warning in hasil["warnings"]:
                    st.warning(f"{hasil['file']}: {warning}")
            except Exception as e:
                st.error(f"Gagal memproses {getattr(file_brs, 'name', 'dokumen')}: {e}")
        st.success(
            f"Migrasi selesai. Total periode yang diproses: Tabel 1 {total_wisman}; "
            f"Tabel 2 {total_hotel}. Periksa peringatan di atas sebelum menggunakan data."
        )


# FUNGSI PERGESERAN TABEL YANG SUDAH DINAMIS

def geser_tabel1(data_lama_negara, data_excel_baru, bulan_sekarang):

    # Ambil 13 kolom sebelumnya (membuang bulan paling lama), lalu tambahkan data bulan ini

    array_baru = data_lama_negara[1:14] + [data_excel_baru]



    # Cari tahu di mana posisi kolom "Jan-Des" berada (Polanya selalu 13 - bulan)

    idx_jandes = 13 - bulan_sekarang



    # Total tahun ini adalah jumlah dari kolom setelah "Jan-Des" sampai akhir

    data_tahun_ini = array_baru[idx_jandes + 1 :]

    total_ini = sum(data_tahun_ini)



    return array_baru + [total_ini]





# --- KONTEN UTAMA (DI TENGAH LAYAR) ---

_, center_col, _ = st.columns([0.5, 6, 0.5])
with center_col:
    # --- HEADER: LOGO & JUDUL BERDAMPINGAN ---
    logo_base64 = get_image_base64("logosibrio.png")
    img_html = f'<img src="data:image/jpeg;base64,{logo_base64}" style="width: 85px; flex-shrink: 0; object-fit: contain;">' if logo_base64 else ''
    html_header = f"""

<div class="header-container">

    {img_html}

    <div class="main-title">SIBRIO</div>

</div>

<div class="main-subtitle">Platform Otomatis untuk Mendukung Penyusunan Berita Resmi Statistik yang<br>Cepat, Tepat, Efisien, dan Terstruktur</div>

"""

    st.markdown(html_header, unsafe_allow_html=True)



    # --- BARIS 1: BULAN & TAHUN ---

    col_b, col_t = st.columns(2)

    with col_b:

        st.markdown('<p class="custom-label">Bulan</p>', unsafe_allow_html=True)

        BULAN = st.selectbox("Bulan", [1,2,3,4,5,6,7,8,9,10,11,12], index=4, label_visibility="collapsed")

    with col_t:

        st.markdown('<p class="custom-label">Tahun</p>', unsafe_allow_html=True)

        TAHUN = st.number_input("Tahun", min_value=2020, max_value=2050, value=2026, label_visibility="collapsed")



    st.write("<br>", unsafe_allow_html=True)



    # --- BARIS 2: TPK YOY ---

    col_tpk1, col_tpk2 = st.columns(2)

    with col_tpk1:

        st.markdown('<p class="custom-label">TPK Bintang bulan yang sama tahun lalu (%)</p>', unsafe_allow_html=True)
        TPK_BINTANG_YOY = st.number_input("TPK Bintang tahun lalu", min_value=0.0, max_value=100.0, value=0.0, format="%.2f", help="Isi TPK bulan yang sama pada tahun sebelumnya, bukan selisih y-o-y.", label_visibility="collapsed")

    with col_tpk2:

        st.markdown('<p class="custom-label">TPK Non-Bintang bulan yang sama tahun lalu (%)</p>', unsafe_allow_html=True)
        TPK_NON_YOY = st.number_input("TPK Non-Bintang tahun lalu", min_value=0.0, max_value=100.0, value=0.0, format="%.2f", help="Isi TPK bulan yang sama pada tahun sebelumnya, bukan selisih y-o-y.", label_visibility="collapsed")



    st.write("<br>", unsafe_allow_html=True)



    # --- BARIS 3: UPLOAD FILE EXCEL & WORD ---

    col_up1, col_up2, col_up3 = st.columns(3)

    with col_up1:

        st.markdown('<p class="custom-label">Data Wisman (Excel)</p>', unsafe_allow_html=True)

        file_wisman = st.file_uploader("Wisman", type=["xlsx", "xls"], label_visibility="collapsed")

    with col_up2:

        st.markdown('<p class="custom-label">Data VHTS (Excel)</p>', unsafe_allow_html=True)

        file_vhts = st.file_uploader("VHTS", type=["xlsx", "xls"], label_visibility="collapsed")

    with col_up3:

        st.markdown('<p class="custom-label">BRS Bulan Lalu (Word)</p>', unsafe_allow_html=True)

        file_brs_lama = st.file_uploader("BRS Lama", type=["docx"], label_visibility="collapsed")



    st.write("<br>", unsafe_allow_html=True)



    # --- TOMBOL GENERATE ---

    if st.button("👇 Generate BRS Sekarang", use_container_width=True):

        PROVINSI = 11 



        if not (0 < float(TPK_BINTANG_YOY) <= 100 and 0 < float(TPK_NON_YOY) <= 100):
            st.error("Isi TPK bulan yang sama tahun sebelumnya untuk hotel bintang dan non-bintang (lebih dari 0 sampai 100).")
        elif file_wisman is not None and file_vhts is not None and file_brs_lama is not None:

            file_template = "Template.docx"




            nama_file_baru = f"BRS_Pariwisata_{BULAN}_{TAHUN}.docx"



            if not os.path.exists(file_template):

                st.error(f"❌ File '{file_template}' tidak ditemukan di folder aplikasi!")

            else:

                with st.spinner('Memproses data dan merangkai dokumen BRS...'):

                    try:

                        # EKSTRAKSI BRS LAMA: SEMUA NILAI DIPETAKAN BERDASARKAN PERIODE
                        doc_lama = docx.Document(file_brs_lama)
                        tahun_sumber = TAHUN - 1 if BULAN == 1 else TAHUN
                        bulan_sumber = 12 if BULAN == 1 else BULAN - 1

                        data_word = {"tabel1": {}, "tabel2": {}}
                        data_word["tabel1"] = extract_tabel1_generation(
                            doc_lama, tahun_sumber, bulan_sumber
                        )

                        # Untuk Tabel 2, ambil dua periode sebelum bulan yang sedang dibuat
                        # berdasarkan label Okt'25/Nov'25/dst., bukan posisi baris Word.
                        tabel2_records, warnings_tabel2 = _parse_tabel2_dari_dokumen(
                            doc_lama, (tahun_sumber, bulan_sumber)
                        )
                        periode_1 = _shift_month(TAHUN, BULAN, -2)
                        periode_2 = _shift_month(TAHUN, BULAN, -1)
                        for periode in (periode_1, periode_2):
                            if periode not in tabel2_records:
                                raise ValueError(
                                    f"Tabel 2 BRS sebelumnya tidak memuat periode "
                                    f"{get_short_bulan(periode[1])} {periode[0]}. "
                                    "Pastikan BRS yang diunggah adalah periode tepat sebelum BRS yang dibuat. "
                                    + ("Peringatan parser: " + " | ".join(warnings_tabel2) if warnings_tabel2 else "")
                                )

                        indikator_tabel2 = {
                            "asing": ("rlm_asing", "rlm_asing_non"),
                            "nusantara": ("rlm_nus", "rlm_nus_non"),
                            "total": ("rlm_total", "rlm_total_non"),
                            "tpk": ("tpk_bintang", "tpk_non"),
                        }
                        for kategori, (key_bintang, key_non) in indikator_tabel2.items():
                            data_word["tabel2"][kategori] = {
                                "baris_lama_1": [
                                    tabel2_records[periode_1][key_bintang],
                                    tabel2_records[periode_1][key_non],
                                ],
                                "baris_lama_2": [
                                    tabel2_records[periode_2][key_bintang],
                                    tabel2_records[periode_2][key_non],
                                ],
                            }

                        # EKSTRAKSI EXCEL WISMAN: CARI KOLOM SIM/SABANG BERDASARKAN HEADER
                        wisman_data = extract_wisman_bulan(file_wisman, BULAN)
                        # BRS lama tetap menjadi sumber utama. Raw bulan-bulan yang tersedia
                        # hanya mengoreksi sel jika terbukti membuat jumlah baris cocok dengan Total.
                        if int(tahun_sumber) == int(TAHUN):
                            corrections_raw = reconcile_tabel1_with_raw_wisman(
                                data_word["tabel1"], file_wisman, TAHUN, bulan_sumber
                            )
                            for correction in corrections_raw:
                                st.warning("Koreksi konsistensi Tabel 1: " + correction)
                        validate_tabel1_totals(data_word["tabel1"], tahun_sumber, bulan_sumber)
                        top10_keys = [
                            "malaysia", "selandia_baru", "perancis", "jerman", "australia",
                            "belanda", "amerika_serikat", "singapura", "inggris", "thailand",
                        ]
                        target_negara = [
                            "Malaysia", "Selandia Baru", "Perancis", "Jerman", "Australia",
                            "Belanda", "Amerika Serikat", "Singapura", "Inggris", "Thailand",
                            "Lainnya", "JUMLAH",
                        ]
                        wisman_sekarang = int(wisman_data["total_wisman"])

                        # EKSTRAKSI EXCEL HOTEL

                        hotel_df = pd.read_excel(file_vhts, sheet_name="Prov_Jenis")

                        hotel_df.columns = hotel_df.columns.str.strip().str.lower()

                        hotel_df["kd_prov"] = pd.to_numeric(hotel_df["kd_prov"], errors="coerce")

                        hotel_df["jenis_akomodasi"] = pd.to_numeric(hotel_df["jenis_akomodasi"], errors="coerce")



                        aceh_bintang = hotel_df[(hotel_df["kd_prov"]==PROVINSI) & (hotel_df["jenis_akomodasi"]==1)]

                        aceh_non = hotel_df[(hotel_df["kd_prov"]==PROVINSI) & (hotel_df["jenis_akomodasi"]==2)]



                        aceh_b = aceh_bintang.iloc[0] if not aceh_bintang.empty else pd.Series()

                        aceh_n = aceh_non.iloc[0] if not aceh_non.empty else pd.Series()



                        hotel_data = {

                            "tpk_bintang": clean_hotel(aceh_b.get("tpk", 0)), "tpk_non": clean_hotel(aceh_n.get("tpk", 0)),

                            "rlm_asing": clean_hotel(aceh_b.get("rlmta", 0)), "rlm_asing_non": clean_hotel(aceh_n.get("rlmta", 0)),

                            "rlm_nus": clean_hotel(aceh_b.get("rlmtnus", 0)), "rlm_nus_non": clean_hotel(aceh_n.get("rlmtnus", 0)),

                            "rlm_total": clean_hotel(aceh_b.get("rlmtgab", 0)), "rlm_total_non": clean_hotel(aceh_n.get("rlmtgab", 0))

                        }

                        # BACKFILL HISTORIS DARI BRS LAMA + SIMPAN BULAN BERJALAN.
                        # BRS lama digunakan sebagai sumber histori yang belum ada.
                        # Setelah itu data Excel bulan berjalan di-UPSERT sebagai data terbaru.
                        # Logika rolling table dan perhitungan BRS tidak diubah.
                        try:
                            brs_bulan = 12 if BULAN == 1 else BULAN - 1
                            brs_tahun = TAHUN - 1 if BULAN == 1 else TAHUN

                            hasil_backfill = backfill_historis_dari_brs(
                                doc_lama, brs_tahun, brs_bulan, data_word
                            )

                            save_historis_wisman(TAHUN, BULAN, wisman_data)
                            save_historis_hotel(TAHUN, BULAN, hotel_data)

                            st.success(
                                "☁️ Historis berhasil diperbarui. "
                                f"Tabel 1: {hasil_backfill['tabel1']} periode dari BRS lama + {get_nama_bulan(BULAN)} {TAHUN}; "
                                f"Tabel 2: {hasil_backfill['tabel2']} periode lama + {get_nama_bulan(BULAN)} {TAHUN}."
                            )
                        except Exception as e_historis:
                            st.warning(
                                "⚠️ BRS berhasil diproses, tetapi proses penyimpanan/backfill historis ke Google Sheets gagal. "
                                f"Periksa konfigurasi Secrets/akses spreadsheet. Detail: {e_historis}"
                            )

                        # KALKULASI NARASI: PERIODE DICARI DENGAN KUNCI (TAHUN, BULAN)
                        data_jumlah = data_word["tabel1"]["JUMLAH"]
                        periode_bulan_lalu = _shift_month(TAHUN, BULAN, -1)
                        periode_yoy = (int(TAHUN) - 1, int(BULAN))

                        wisman_tahun_lalu = period_value(data_jumlah, periode_yoy, "Wisman")
                        wisman_bulan_lalu = period_value(data_jumlah, periode_bulan_lalu, "Wisman")

                        # Kumulatif tahun berjalan: bulan-bulan sebelum target dari BRS lama + target dari Excel.
                        wisman_kumulatif_ini = int(wisman_sekarang)
                        for m in range(1, int(BULAN)):
                            wisman_kumulatif_ini += period_value(
                                data_jumlah, (int(TAHUN), m), "Wisman kumulatif berjalan"
                            )

                        # Kumulatif pembanding Jan s.d. bulan yang sama tahun lalu.
                        if int(tahun_sumber) == int(TAHUN):
                            kumulatif_tahun_lalu_penuh = int(data_jumlah["cumulative"])
                            bulan_setelah_target = range(int(BULAN) + 1, 13)
                            wisman_kumulatif_lalu = kumulatif_tahun_lalu_penuh - sum(
                                period_value(data_jumlah, (int(TAHUN) - 1, m), "Wisman kumulatif tahun lalu")
                                for m in bulan_setelah_target
                            )
                        else:
                            wisman_kumulatif_lalu = sum(
                                period_value(data_jumlah, (int(TAHUN) - 1, m), "Wisman kumulatif tahun lalu")
                                for m in range(1, 13)
                            )

                        mtm_w = ((wisman_sekarang - wisman_bulan_lalu) / wisman_bulan_lalu) * 100 if wisman_bulan_lalu else 0
                        yoy_w = ((wisman_sekarang - wisman_tahun_lalu) / wisman_tahun_lalu) * 100 if wisman_tahun_lalu else 0
                        ctc_w = ((wisman_kumulatif_ini - wisman_kumulatif_lalu) / wisman_kumulatif_lalu) * 100 if wisman_kumulatif_lalu else 0

                        tpk_bintang_lalu = data_word["tabel2"]["tpk"]["baris_lama_2"][0]
                        tpk_non_lalu = data_word["tabel2"]["tpk"]["baris_lama_2"][1]

                        # SETUP CONTEXT

                        bln_lalu_angka = 12 if BULAN == 1 else BULAN - 1

                        thn_mtm = TAHUN - 1 if BULAN == 1 else TAHUN

                        bln_singkat = [

                            f"{get_short_bulan(bln_lalu_angka - 1 if bln_lalu_angka > 1 else 12)}'{str(thn_mtm - 1 if bln_lalu_angka == 1 else thn_mtm)[-2:]}",

                            f"{get_short_bulan(bln_lalu_angka)}'{str(thn_mtm)[-2:]}", f"{get_short_bulan(BULAN)}'{str(TAHUN)[-2:]}"

                        ]



                        context = {

                            "bulan": get_nama_bulan(BULAN), "tahun": str(TAHUN),

                            "bulan_lalu": get_nama_bulan(bln_lalu_angka), "tahun_lalu": str(TAHUN - 1), "tahun_mtm": str(thn_mtm),

                            "total_wisman": format_ribuan(wisman_sekarang), "total_wisman_lalu": format_ribuan(wisman_bulan_lalu),

                            "total_wisman_yoy": format_ribuan(wisman_tahun_lalu), "kumulatif_wisman": format_ribuan(wisman_kumulatif_ini),

                            "kumulatif_wisman_lalu": format_ribuan(wisman_kumulatif_lalu),

                            "mtm_wisman": format_desimal(abs(mtm_w)), "trend_mtm_wisman": get_arah_tren(mtm_w),

                            "yoy_wisman": format_desimal(abs(yoy_w)), "arah_yoy_wisman": get_arah_tren_singkat(yoy_w),

                            "trend_yoy_wisman": get_arah_tren(yoy_w), "ctc_wisman": format_desimal(abs(ctc_w)),

                            "trend_ctc_wisman": get_arah_tren(ctc_w), "negara_dominan": "Malaysia",

                            "persen_negara_dominan": format_desimal((wisman_data.get("malaysia", 0) / wisman_sekarang) * 100) if wisman_sekarang else "0",

                            "tpk_bintang": format_desimal(hotel_data.get("tpk_bintang", 0)), "tpk_bintang_lalu": format_desimal(tpk_bintang_lalu),

                            "selisih_tpk_bintang": format_desimal(abs(hotel_data.get("tpk_bintang", 0) - tpk_bintang_lalu)), "arah_tpk_bintang": get_arah_tren_singkat(hotel_data.get("tpk_bintang", 0) - tpk_bintang_lalu),

                            "trend_tpk_bintang_mtm": get_arah_tren(hotel_data.get("tpk_bintang", 0) - tpk_bintang_lalu),

                            "tpk_bintang_yoy": format_desimal(TPK_BINTANG_YOY), "selisih_tpk_bintang_yoy": format_desimal(abs(hotel_data.get("tpk_bintang", 0) - TPK_BINTANG_YOY)),

                            "trend_tpk_bintang_yoy": get_arah_tren(hotel_data.get("tpk_bintang", 0) - TPK_BINTANG_YOY),

                            "tpk_non": format_desimal(hotel_data.get("tpk_non", 0)), "selisih_tpk_non_mtm": format_desimal(abs(hotel_data.get("tpk_non", 0) - tpk_non_lalu)),

                            "trend_tpk_non_mtm": get_arah_tren(hotel_data.get("tpk_non", 0) - tpk_non_lalu), "selisih_tpk_non_yoy": format_desimal(abs(hotel_data.get("tpk_non", 0) - TPK_NON_YOY)),

                            "trend_tpk_non_yoy": get_arah_tren(hotel_data.get("tpk_non", 0) - TPK_NON_YOY),

                            "rlm_total": format_desimal(hotel_data.get("rlm_total", 0)), "rlm_total_lalu": format_desimal(data_word["tabel2"]["total"]["baris_lama_2"][0]),

                            "selisih_rlm_total_mtm": format_desimal(abs(hotel_data.get("rlm_total", 0) - data_word["tabel2"]["total"]["baris_lama_2"][0])), "trend_rlm_total_mtm": get_arah_tren(hotel_data.get("rlm_total", 0) - data_word["tabel2"]["total"]["baris_lama_2"][0]),

                            "rlm_asing": format_desimal(hotel_data.get("rlm_asing", 0)), "rlm_asing_lalu": format_desimal(data_word["tabel2"]["asing"]["baris_lama_2"][0]),

                            "selisih_rlm_asing_mtm": format_desimal(abs(hotel_data.get("rlm_asing", 0) - data_word["tabel2"]["asing"]["baris_lama_2"][0])), "trend_rlm_asing_mtm": get_arah_tren(hotel_data.get("rlm_asing", 0) - data_word["tabel2"]["asing"]["baris_lama_2"][0]),

                            "rlm_nus": format_desimal(hotel_data.get("rlm_nus", 0)), "selisih_rlm_nus_mtm": format_desimal(abs(hotel_data.get("rlm_nus", 0) - data_word["tabel2"]["nusantara"]["baris_lama_2"][0])),

                            "trend_rlm_nus_mtm": get_arah_tren(hotel_data.get("rlm_nus", 0) - data_word["tabel2"]["nusantara"]["baris_lama_2"][0]),



                            "bln_1": bln_singkat[0], "bln_2": bln_singkat[1], "bln_3": bln_singkat[2],

                            "as_b1": format_desimal(data_word["tabel2"]["asing"]["baris_lama_1"][0]), "as_n1": format_desimal(data_word["tabel2"]["asing"]["baris_lama_1"][1]),

                            "as_b2": format_desimal(data_word["tabel2"]["asing"]["baris_lama_2"][0]), "as_n2": format_desimal(data_word["tabel2"]["asing"]["baris_lama_2"][1]),

                            "as_b3": format_desimal(hotel_data.get("rlm_asing", 0)), "as_n3": format_desimal(hotel_data.get("rlm_asing_non", 0)),

                            "nus_b1": format_desimal(data_word["tabel2"]["nusantara"]["baris_lama_1"][0]), "nus_n1": format_desimal(data_word["tabel2"]["nusantara"]["baris_lama_1"][1]),

                            "nus_b2": format_desimal(data_word["tabel2"]["nusantara"]["baris_lama_2"][0]), "nus_n2": format_desimal(data_word["tabel2"]["nusantara"]["baris_lama_2"][1]),

                            "nus_b3": format_desimal(hotel_data.get("rlm_nus", 0)), "nus_n3": format_desimal(hotel_data.get("rlm_nus_non", 0)),

                            "tot_b1": format_desimal(data_word["tabel2"]["total"]["baris_lama_1"][0]), "tot_n1": format_desimal(data_word["tabel2"]["total"]["baris_lama_1"][1]),

                            "tot_b2": format_desimal(data_word["tabel2"]["total"]["baris_lama_2"][0]), "tot_n2": format_desimal(data_word["tabel2"]["total"]["baris_lama_2"][1]),

                            "tot_b3": format_desimal(hotel_data.get("rlm_total", 0)), "tot_n3": format_desimal(hotel_data.get("rlm_total_non", 0)),

                            "tpk_b1": format_desimal(data_word["tabel2"]["tpk"]["baris_lama_1"][0]), "tpk_n1": format_desimal(data_word["tabel2"]["tpk"]["baris_lama_1"][1]),

                            "tpk_b2": format_desimal(data_word["tabel2"]["tpk"]["baris_lama_2"][0]), "tpk_n2": format_desimal(data_word["tabel2"]["tpk"]["baris_lama_2"][1]),

                            "tpk_b3": format_desimal(hotel_data.get("tpk_bintang", 0)), "tpk_n3": format_desimal(hotel_data.get("tpk_non", 0)),

                        }



                        # RENDER KE WORD

                        doc_tpl = DocxTemplate(file_template)

                        doc_tpl.render(context)

                        docx_obj = doc_tpl.docx 



                        tabel1_word = docx_obj.tables[0]

                        nama_bulan_s = ["Jan", "Feb", "Mar", "Apr", "Mei", "Jun", "Jul", "Ags", "Sep", "Okt", "Nov", "Des"]

                        headers_tabel1 = nama_bulan_s[BULAN - 1:] + ["Jan-Des"] + nama_bulan_s[:BULAN] + ["Total"]

                        row_header = tabel1_word.rows[1].cells[-15:]
                        for col_idx in range(15):
                            row_header[col_idx].text = headers_tabel1[col_idx]
                        update_tabel1_year_groups(tabel1_word, TAHUN, BULAN)



                        baris_idx = 3

                        for key, negara in zip(top10_keys + ["lainnya"], target_negara[:-1]):
                            baris_baru = build_tabel1_output_row(
                                data_word["tabel1"][negara], wisman_data[key],
                                TAHUN, BULAN, tahun_sumber
                            )
                            sel_tulis = tabel1_word.rows[baris_idx].cells[-15:]
                            for col_idx in range(15):
                                sel_tulis[col_idx].text = format_ribuan(baris_baru[col_idx])
                            baris_idx += 1

                        baris_baru_jumlah = build_tabel1_output_row(
                            data_word["tabel1"]["JUMLAH"], wisman_data["total_wisman"],
                            TAHUN, BULAN, tahun_sumber
                        )
                        sel_tulis = tabel1_word.rows[baris_idx].cells[-15:]
                        for col_idx in range(15):
                            sel_tulis[col_idx].text = format_ribuan(baris_baru_jumlah[col_idx])



                        # SIMPAN KE MEMORY UNTUK FITUR DOWNLOAD BUTTON

                        output_stream = io.BytesIO()

                        doc_tpl.save(output_stream)

                        output_stream.seek(0)



                        st.success(f"🎉 Berhasil! Silakan download dokumen BRS Anda di bawah ini:")

                        st.download_button(

                            label=f"📥 Download {nama_file_baru}",

                            data=output_stream,

                            file_name=nama_file_baru,

                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",

                            use_container_width=True

                        )



                    except Exception as e:

                        st.error(f"Terjadi kesalahan saat memproses data: {e}")

        else:

            st.warning("⚠️ Harap lengkapi ketiga file (Wisman, VHTS, dan BRS Lama) terlebih dahulu!")

# --- FITUR HISTORIS DATA ---
# Diletakkan di luar blok Generate agar selalu dapat digunakan setelah database terisi.
tampilkan_historis_data()
tampilkan_migrasi_historis()
